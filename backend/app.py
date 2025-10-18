import json
import os
import shutil
import subprocess
import uuid
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS
from flask_socketio import SocketIO, emit, join_room, leave_room
from pptx import Presentation

BASE_DIR = Path(__file__).resolve().parent
STORAGE_DIR = BASE_DIR / "storage"
DATA_FILE = STORAGE_DIR / "data" / "classes.json"
UPLOAD_DIR = STORAGE_DIR / "uploads"
CONVERSION_DIR = STORAGE_DIR / "conversions"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
CONVERSION_DIR.mkdir(parents=True, exist_ok=True)
DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
if not DATA_FILE.exists():
    DATA_FILE.write_text(json.dumps({"classes": []}))

# Load environment variables from .env (if present). In production prefer real env vars or secrets manager.
load_dotenv()

# Example secret/config variables (read but not printed). Keep these on the server only.
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
FRONTEND_ORIGIN = os.getenv("FRONTEND_ORIGIN", "http://localhost:5173")

# Telegram Bot (optional)
telegram_bot_instance = None


def create_app() -> Flask:
    app = Flask(__name__)
    app.config["SECRET_KEY"] = "class-assist-secret"
    app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB
    CORS(
        app,
        resources={r"/api/*": {"origins": [FRONTEND_ORIGIN]}},
        supports_credentials=True,
    )
    return app


def create_socketio(flask_app: Flask) -> SocketIO:
    return SocketIO(
        flask_app,
        cors_allowed_origins="*",
        async_mode="eventlet",
    )


app = create_app()
socketio = create_socketio(app)


class StorageError(RuntimeError):
    """Custom exception for storage operations."""


class ClassStore:
    def __init__(self, data_file: Path) -> None:
        self.data_file = data_file

    def _read(self) -> dict:
        try:
            raw = self.data_file.read_text()
            return json.loads(raw)
        except json.JSONDecodeError as exc:
            raise StorageError("Invalid data store format") from exc

    def _write(self, data: dict) -> None:
        tmp_path = self.data_file.with_suffix(".tmp")
        tmp_path.write_text(json.dumps(data, indent=2))
        tmp_path.replace(self.data_file)

    def list_classes(self) -> list:
        classes = self._read().get("classes", [])
        return [self._hydrate_class(class_data) for class_data in classes]

    def save_class(self, class_data: dict) -> dict:
        data = self._read()
        classes = data.setdefault("classes", [])
        existing_index = next((i for i, c in enumerate(classes) if c["id"] == class_data["id"]), None)
        if existing_index is None:
            classes.append(class_data)
        else:
            classes[existing_index] = class_data
        self._write(data)
        return class_data

    def get_class(self, class_id: str) -> dict | None:
        class_record = next((c for c in self._read().get("classes", []) if c["id"] == class_id), None)
        if not class_record:
            return None
        hydrated = self._hydrate_class(class_record)
        if hydrated is not class_record:
            self.save_class(hydrated)
        return hydrated

    def _hydrate_class(self, class_record: dict) -> dict:
        total_slides = class_record.get("total_slides")
        if isinstance(total_slides, int) and total_slides > 0:
            return class_record
        updated = dict(class_record)
        slides = updated.get("slides") or []
        pdf_path = CONVERSION_DIR / f"{updated['id']}.pdf"
        updated["total_slides"] = len(slides) if slides else count_pdf_pages(pdf_path)
        return updated


store = ClassStore(DATA_FILE)


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in {"pptx", "ppt", "pdf"}


def convert_to_pdf(source_path: Path, target_path: Path) -> None:
    """Convert presentation to PDF using LibreOffice headless mode."""
    target_path.parent.mkdir(parents=True, exist_ok=True)
    if source_path.suffix.lower() == ".pdf":
        shutil.copy(source_path, target_path)
        return
    command = [
        "soffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(target_path.parent),
        str(source_path),
    ]
    run = subprocess.run(command, capture_output=True, text=True)
    if run.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {run.stderr}")
    converted_path = target_path.parent / f"{source_path.stem}.pdf"
    if converted_path != target_path:
        converted_path.replace(target_path)


def extract_slide_text(source_path: Path) -> list[dict]:
    """Extract basic text content from each slide using python-pptx."""
    if source_path.suffix.lower() == ".pdf":
        return []
    presentation = Presentation(str(source_path))
    slides_summary: list[dict] = []
    for index, slide in enumerate(presentation.slides):
        text_runs: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    text_runs.append(text)
        slides_summary.append({
            "index": index,
            "text": "\n".join(text_runs),
        })
    return slides_summary


def count_pdf_pages(pdf_path: Path) -> int:
    try:
        data = pdf_path.read_bytes()
    except OSError:
        return 0
    count = data.count(b"/Type /Page")
    return count or 1


@app.route("/api/health", methods=["GET"])
def health() -> tuple[dict, int]:
    return jsonify({"status": "ok"}), 200


# Lightweight request logging to help debug unexpected 404s from clients
@app.before_request
def log_request_info() -> None:
    # Log method and path for every incoming request to aid tracing
    app.logger.info("Incoming request: %s %s", request.method, request.path)


@app.route("/api/classes", methods=["GET"])
def list_classes() -> tuple[dict, int]:
    return jsonify({"classes": store.list_classes()}), 200


@app.route("/api/classes", methods=["POST"])
def create_class() -> tuple[dict, int]:
    if "file" not in request.files:
        return jsonify({"error": "file is required"}), 400

    file_storage = request.files["file"]
    if file_storage.filename == "":
        return jsonify({"error": "file must have a name"}), 400
    if not allowed_file(file_storage.filename):
        return jsonify({"error": "unsupported file type"}), 400

    class_id = str(uuid.uuid4())
    upload_name = f"{class_id}_{file_storage.filename}"
    uploaded_path = UPLOAD_DIR / upload_name
    file_storage.save(uploaded_path)

    pdf_target = CONVERSION_DIR / f"{class_id}.pdf"
    try:
        convert_to_pdf(uploaded_path, pdf_target)
    except RuntimeError as exc:
        uploaded_path.unlink(missing_ok=True)
        return jsonify({"error": str(exc)}), 500

    slides_summary = extract_slide_text(uploaded_path)
    total_slides = len(slides_summary) if slides_summary else count_pdf_pages(pdf_target)

    class_record = {
        "id": class_id,
        "title": request.form.get("title") or file_storage.filename,
        "course": request.form.get("course") or "",
        "level": request.form.get("level") or "",
        "created_at": datetime.utcnow().isoformat() + "Z",
        "presentation_file": upload_name,
        "pdf_path": f"/api/classes/{class_id}/pdf",
        "slides": slides_summary,
        "total_slides": total_slides,
        "status": "ready",
        "current_slide": 0,
        "session": None,
    }
    store.save_class(class_record)
    return jsonify(class_record), 201


@app.route("/api/classes/<class_id>", methods=["GET"])
def get_class(class_id: str) -> tuple[dict, int]:
    class_record = store.get_class(class_id)
    if not class_record:
        return jsonify({"error": "class not found"}), 404
    return jsonify(class_record), 200


@app.route("/api/classes/<class_id>/pdf", methods=["GET"])
def get_class_pdf(class_id: str):
    pdf_path = CONVERSION_DIR / f"{class_id}.pdf"
    if not pdf_path.exists():
        return jsonify({"error": "pdf not found"}), 404
    return send_from_directory(CONVERSION_DIR, pdf_path.name)


# Compatibility endpoints: some older frontend builds (or external tools)
# may call /api/sessions/active or /api/presentations. Provide small
# compatible responses so those clients don't see 404s while we
# identify and fix the root cause in the front-end.
@app.route("/api/sessions/active", methods=["GET"])
def sessions_active() -> tuple[dict, int]:
    classes = store.list_classes()
    active = [c for c in classes if c.get("status") == "live"]
    return jsonify({"active": active}), 200


@app.route("/api/presentations", methods=["GET"])
def list_presentations() -> tuple[dict, int]:
    classes = store.list_classes()
    presentations = [
        {"id": c["id"], "title": c.get("title"), "pdf_path": c.get("pdf_path")} for c in classes
    ]
    return jsonify({"presentations": presentations}), 200


@app.route("/api/classes/<class_id>/start", methods=["POST"])
def start_class(class_id: str) -> tuple[dict, int]:
    class_record = store.get_class(class_id)
    if not class_record:
        return jsonify({"error": "class not found"}), 404
    class_record["status"] = "live"
    class_record["session"] = {
        "started_at": datetime.utcnow().isoformat() + "Z",
    }
    class_record["current_slide"] = 0
    store.save_class(class_record)
    socketio.emit("class_started", {"classId": class_id}, room=class_id)
    return jsonify(class_record), 200


@app.route("/api/classes/<class_id>/stop", methods=["POST"])
def stop_class(class_id: str) -> tuple[dict, int]:
    class_record = store.get_class(class_id)
    if not class_record:
        return jsonify({"error": "class not found"}), 404
    class_record["status"] = "ended"
    store.save_class(class_record)
    socketio.emit("class_stopped", {"classId": class_id}, room=class_id)
    return jsonify(class_record), 200


def _update_slide(class_id: str, new_index: int) -> dict | None:
    class_record = store.get_class(class_id)
    if not class_record:
        return None
    class_record["current_slide"] = new_index
    store.save_class(class_record)
    return class_record


def process_control_action(class_id: str, action: str, index: int | None = None) -> dict | None:
    """
    Procesa una acción de control (next, prev, goto) y actualiza la diapositiva.
    Retorna el class_record actualizado o None si hay error.
    Esta función puede ser llamada desde WebSocket handlers o desde el bot de Telegram.
    """
    class_record = store.get_class(class_id)
    if not class_record:
        return None

    if action not in {"next", "prev", "goto"}:
        return None

    current = class_record.get("current_slide", 0)
    slides_total = (
        class_record.get("total_slides")
        or len(class_record.get("slides", []))
        or _count_pdf_slides(class_id)
    )

    if action == "next":
        current = min(current + 1, max(slides_total - 1, 0))
    elif action == "prev":
        current = max(current - 1, 0)
    elif action == "goto":
        if isinstance(index, int):
            current = max(0, min(index, max(slides_total - 1, 0)))

    updated = _update_slide(class_id, current)
    if updated:
        # Emitir evento a todos los clientes conectados a esta clase
        socketio.emit(
            "slide_update",
            {"classId": class_id, "currentSlide": updated["current_slide"]},
            room=class_id,
        )
    return updated


@socketio.on("connect")
def handle_connect():
    emit("connected", {"message": "socket connected"})


@socketio.on("join_class")
def handle_join(data):
    class_id = data.get("classId")
    if not class_id:
        emit("error", {"error": "classId is required"})
        return
    join_room(class_id)
    emit("joined", {"classId": class_id})
    class_record = store.get_class(class_id)
    if class_record:
        emit("slide_update", {"classId": class_id, "currentSlide": class_record["current_slide"]})


@socketio.on("leave_class")
def handle_leave(data):
    class_id = data.get("classId")
    if class_id:
        leave_room(class_id)
        emit("left", {"classId": class_id})


@socketio.on("control_action")
def handle_control_action(data):
    class_id = data.get("classId")
    action = data.get("action")
    index = data.get("index")
    
    if not class_id or action not in {"next", "prev", "goto"}:
        emit("error", {"error": "invalid control action"})
        return

    updated = process_control_action(class_id, action, index)
    if updated is None:
        emit("error", {"error": "unable to update slide"})


def _count_pdf_slides(class_id: str) -> int:
    pdf_path = CONVERSION_DIR / f"{class_id}.pdf"
    if not pdf_path.exists():
        return 0
    return count_pdf_pages(pdf_path)


async def start_telegram_bot():
    """Inicia el bot de Telegram si el token está configurado."""
    global telegram_bot_instance
    if not TELEGRAM_BOT_TOKEN:
        app.logger.warning("TELEGRAM_BOT_TOKEN not set. Telegram bot will not start.")
        return

    try:
        from telegram_bot import TelegramBot
        telegram_bot_instance = TelegramBot(
            TELEGRAM_BOT_TOKEN, 
            socketio, 
            store,
            process_control_action
        )
        await telegram_bot_instance.start_bot()
        app.logger.info("Telegram bot initialized successfully")
    except Exception as exc:
        app.logger.error(f"Failed to start Telegram bot: {exc}")


if __name__ == "__main__":
    import asyncio
    
    # Iniciar bot de Telegram en segundo plano si está configurado
    if TELEGRAM_BOT_TOKEN:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        loop.create_task(start_telegram_bot())
        
        # Correr eventloop en background thread para no bloquear Flask
        import threading
        def run_async_loop():
            loop.run_forever()
        threading.Thread(target=run_async_loop, daemon=True).start()
    
    socketio.run(app, host="0.0.0.0", port=5000)
